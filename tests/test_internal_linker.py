import csv
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import internal_linker as il  # noqa: E402

NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def write_map(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def test_load_rules_skips_header_and_dedupes(tmp_path):
    map_path = tmp_path / "map.csv"
    write_map(map_path, [
        ["Keyword", "URL"],
        ["seo", "https://example.com/seo"],
        ["seo", "https://example.com/duplicate-ignored"],
        ["", "https://example.com/missing-phrase"],
    ])
    rules = il.load_rules(map_path)
    assert len(rules) == 1
    assert rules[0].phrase == "seo"
    assert rules[0].url == "https://example.com/seo"


def test_load_rules_per_row_occurrence_cap(tmp_path):
    map_path = tmp_path / "map.csv"
    write_map(map_path, [["internal linking", "/x", "3"]])
    rules = il.load_rules(map_path)
    assert rules[0].occurrence_cap == 3


def test_single_run_phrase_linked_once_only(tmp_path):
    doc = Document()
    doc.add_paragraph("We care about seo and seo and seo.")
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [["seo", "https://example.com/seo"]])
    engine = il.LinkEngine(il.load_rules(map_path))

    document = Document(docx_path)
    total = il.link_document(document, engine, link_color="0563C1", underline_links=True)
    assert total == 1
    para = document.paragraphs[0]
    assert para.text == "We care about seo and seo and seo."
    assert len(para._p.findall(f".//{NS}hyperlink")) == 1


def test_phrase_split_across_runs_is_still_found(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Learn about ")
    p.add_run("knowled")
    p.add_run("ge graph").italic = True
    p.add_run(" basics.")
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [["knowledge graph", "https://example.com/kg"]])
    engine = il.LinkEngine(il.load_rules(map_path))

    document = Document(docx_path)
    total = il.link_document(document, engine, link_color="0563C1", underline_links=True)
    assert total == 1
    assert document.paragraphs[0].text == "Learn about knowledge graph basics."


def test_longer_phrase_wins_on_overlap(tmp_path):
    doc = Document()
    doc.add_paragraph("We publish data analysis reports.")
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [
        ["data", "https://example.com/data"],
        ["data analysis", "https://example.com/data-analysis"],
    ])
    engine = il.LinkEngine(il.load_rules(map_path))

    document = Document(docx_path)
    il.link_document(document, engine, link_color=None, underline_links=False)
    hyperlink_texts = [t.text for t in document.paragraphs[0]._p.findall(f".//{NS}hyperlink//{NS}t")]
    assert hyperlink_texts == ["data analysis"]


def test_text_inside_an_existing_hyperlink_is_never_touched_again(tmp_path):
    doc = Document()
    doc.add_paragraph("We care about seo basics.")
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [["seo", "https://example.com/seo"]])
    rules = il.load_rules(map_path)
    engine = il.LinkEngine(rules)

    document = Document(docx_path)
    total = il.link_document(document, engine, link_color="0563C1", underline_links=True)
    assert total == 1

    # Re-run against the now-linked paragraph: the word lives inside a
    # <w:hyperlink> now, which paragraph.runs does not walk, so a second
    # pass (even after resetting the cap) must find nothing to do.
    engine.reset()
    again = il.link_paragraph(document.paragraphs[0], engine, link_color="0563C1", underline_links=True)
    assert again == 0


def test_occurrence_cap_spans_the_whole_document(tmp_path):
    doc = Document()
    doc.add_paragraph("internal linking helps seo.")
    doc.add_paragraph("more on internal linking later.")
    doc.add_paragraph("and once more: internal linking.")
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [["internal linking", "/guides/internal-linking", "2"]])
    engine = il.LinkEngine(il.load_rules(map_path))

    document = Document(docx_path)
    total = il.link_document(document, engine, link_color="0563C1", underline_links=True)
    assert total == 2


def test_table_cells_including_nested_are_walked(tmp_path):
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.rows[0].cells[0].text = "internal linking in an outer cell"
    inner = outer.rows[0].cells[0].add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = "internal linking in a nested cell"
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [["internal linking", "/guides/internal-linking", "5"]])
    engine = il.LinkEngine(il.load_rules(map_path))

    document = Document(docx_path)
    total = il.link_document(document, engine, link_color="0563C1", underline_links=True)
    assert total == 2


def test_formatting_is_preserved_on_both_sides_of_a_link(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Before ").bold = True
    p.add_run("seo")
    run_after = p.add_run(" after")
    run_after.italic = True
    docx_path = tmp_path / "doc.docx"
    doc.save(docx_path)

    map_path = tmp_path / "map.csv"
    write_map(map_path, [["seo", "https://example.com/seo"]])
    engine = il.LinkEngine(il.load_rules(map_path))

    document = Document(docx_path)
    il.link_document(document, engine, link_color="0563C1", underline_links=True)
    runs = document.paragraphs[0].runs
    assert runs[0].bold is True
    assert runs[0].text == "Before "
    assert runs[-1].italic is True
    assert runs[-1].text == " after"


def test_discover_docx_files_batch_and_recursive(tmp_path):
    (tmp_path / "sub").mkdir()
    Document().save(tmp_path / "a.docx")
    Document().save(tmp_path / "sub" / "b.docx")
    (tmp_path / "~$a.docx").write_text("lock file")  # Word lock files must be ignored

    flat = il.discover_docx_files(tmp_path, recursive=False)
    assert [p.name for p in flat] == ["a.docx"]

    recursive = il.discover_docx_files(tmp_path, recursive=True)
    assert {p.name for p in recursive} == {"a.docx", "b.docx"}


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
